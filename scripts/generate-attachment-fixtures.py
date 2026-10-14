"""Create small, known-content files for attachment extraction checks."""
from pathlib import Path
from reportlab.pdfgen import canvas
from docx import Document
from pptx import Presentation
from pptx.util import Inches
from PIL import Image, ImageDraw

root = Path(__file__).parent / "fixtures" / "attachments"
root.mkdir(parents=True, exist_ok=True)

pdf = canvas.Canvas(str(root / "sample.pdf"))
for page in range(1, 4):
    pdf.setFont("Helvetica", 14)
    pdf.drawString(60, 760, f"Page {page}: document extraction sample")
    if page == 3:
        pdf.drawString(60, 730, "SOFTMAX_STABILITY_ANCHOR is on page three.")
    pdf.showPage()
pdf.save()

scan_image = Image.new("RGB", (900, 220), "white")
ImageDraw.Draw(scan_image).text((30, 70), "SCANNED_PAGE_ANCHOR is only in pixels", fill="black")
scan_image.save(root / "sample.png")
scan = canvas.Canvas(str(root / "scanned.pdf"))
scan.drawInlineImage(scan_image, 50, 550, width=500, height=125)
scan.showPage()
scan.save()

doc = Document()
doc.add_heading("Word 样本", 0)
doc.add_paragraph("DOCX_ANCHOR：导入后应保留这段文字。")
doc.save(root / "sample.docx")

deck = Presentation()
slide = deck.slides.add_slide(deck.slide_layouts[5])
box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(7), Inches(2))
box.text = "PPTX_ANCHOR：这是幻灯片上的文字。"
deck.save(root / "sample.pptx")
